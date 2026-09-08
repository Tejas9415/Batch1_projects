import os
from docx import Document

from core.state import GraphState

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)),"generate_resumes")

def docx_writer_node(state: GraphState) -> dict:
    os.makedirs(OUTPUT_DIR,exist_ok=True)

    tailored = state["tailored_resume"]
    doc = Document()

    doc.add_heading(f"{state.get('role_title','Tailored Resume')} - {state.get('company_name','')}", level=1)

    doc.add_heading("Professional_Summary", level=2)
    doc.add_paragraph(tailored.professional_summary)

    doc.add_heading("Skills",level=2)
    doc.add_paragraph(", ".join(tailored.skills_section))

    doc.add_heading("Experience", level=2)
    for bullet in tailored.experience_bullets:
        doc.add_paragraph(bullet, style="List Bullet")
    
    name = state.get('company_name','company').replace(" ","_") #Maruti Suzuki
    filename = f"tailored_resume_{name}.docx"
    path = os.path.join(OUTPUT_DIR, filename)
    doc.save(path) #doc got saved on the disk

    return {
        "tailored_resume_docx_path": path,
        "messages": [
            {
                "role":"assistant",
                "content":f"[DOCX] saved to {path}"
            }
        ]
    }